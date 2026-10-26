"""The whole report pipeline on made-up portal exports, through the service."""

import io

import docx
import make_report
import pdf_export
import pytest
import ritm_fetch
import synthetic
from department_info import DepartmentInfo
from helpers import FakePortal, finished, log_in, make_service


def offline_reports(request, workdir, log, pdf=False):
    # PDF takes LibreOffice seconds per document; only the tests about PDF ask for it
    synthetic.write_attestation(workdir)
    department = DepartmentInfo(request.department, request.faculty_title, request.head_name or "И. О. Фамилия")
    period = ritm_fetch.AttestationPeriod(2026, True, 1)
    return make_report.build_all(None, period, workdir, ["б", "м"], log=log, department=department, pdf=pdf)


def finished_job(tmp_path, build=offline_reports, **overrides):
    _, client, _ = make_service(tmp_path, FakePortal(build=build))
    log_in(client)
    return client, finished(client, **overrides)


def download(job, label: str, kind: str = "DOCX") -> dict:
    document = next(item for item in job["downloads"] if item["label"] == label)
    return next(item for item in document["files"] if item["format"] == kind)


def short_report_text(client, job) -> str:
    short = download(job, "Краткий отчёт: бакалавриат")
    document = docx.Document(io.BytesIO(client.get(short["href"]).content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs) + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )


def test_cannot_lose_any_report_file(tmp_path):
    _, job = finished_job(tmp_path)
    assert len(job["downloads"]) == 5, f"Expected two short, two detailed and Markdown reports: {job['downloads']}"


def test_cannot_drop_master_group_from_markdown(tmp_path):
    _, job = finished_job(tmp_path)
    assert "ИТСм(о)-11" in job["markdown"], "Master's group is missing from the Markdown report"


def test_cannot_download_short_report_as_something_else(tmp_path):
    client, job = finished_job(tmp_path)
    short = download(job, "Краткий отчёт: магистратура")
    assert client.get(short["href"]).content[:2] == b"PK", "Short report download is not a DOCX file"


def test_cannot_show_department_rank_wrong(tmp_path):
    _, job = finished_job(tmp_path)
    assert "место кафедры: **2 из 3**" in job["markdown"], "Department rank in the summary is wrong"


def test_cannot_print_empty_critical_table(tmp_path):
    _, job = finished_job(tmp_path)
    assert "Таких дисциплин нет." in job["markdown"], "Empty critical-disciplines table was printed as a bare header"


def test_cannot_title_other_department_as_itrs(tmp_path):
    _, job = finished_job(tmp_path, department="ВМ", faculty_title="Факультет информатики")
    assert job["markdown"].startswith("# Кафедра ВМ:"), "Report of another department is titled ИТРС"


def test_cannot_rank_other_department_as_itrs(tmp_path):
    _, job = finished_job(tmp_path, department="ВМ", faculty_title="Факультет информатики")
    assert "место кафедры: **3 из 3**" in job["markdown"], "Rank of another department came from ИТРС"


def test_cannot_put_wrong_faculty_in_short_report_header(tmp_path):
    client, job = finished_job(tmp_path, department="ВМ", faculty_title="Факультет информатики")
    assert "Факультет информатики" in short_report_text(client, job), "Short report header has a wrong faculty"


def test_cannot_sign_short_report_by_wrong_head(tmp_path):
    client, job = finished_job(tmp_path, head_name="П. П. Петров")
    assert "П. П. Петров" in short_report_text(client, job), "Short report is signed by a wrong head"


def test_cannot_send_page_without_chart_data(tmp_path):
    _, job = finished_job(tmp_path)
    assert len(job["data"]["groups"]) == 3, "Job result has no chart data for the page"


def pdf_reports(request, workdir, log):
    return offline_reports(request, workdir, log, pdf=True)


@pytest.mark.skipif(pdf_export.find_soffice() is None, reason="LibreOffice is not installed")
def test_cannot_lose_pdf_of_short_report(tmp_path):
    client, job = finished_job(tmp_path, build=pdf_reports)
    pdf = download(job, "Краткий отчёт: бакалавриат", "PDF")
    assert client.get(pdf["href"]).content[:5] == b"%PDF-", "Short report has no PDF download"


def test_cannot_fail_report_when_libreoffice_is_missing(tmp_path, monkeypatch):
    monkeypatch.setattr(pdf_export, "find_soffice", lambda: None)
    _, job = finished_job(tmp_path, build=pdf_reports)
    formats = {file["format"] for item in job["downloads"] for file in item["files"]}
    assert formats == {"DOCX", "Markdown"}, f"Without LibreOffice the reports must stay in DOCX: {formats}"


def offline_group_reports(request, workdir, log):
    attestation_dir = synthetic.write_attestation(workdir)
    # what fetch_level leaves behind for chosen groups: only their exports
    for path in attestation_dir.glob("[бм]/*.xml"):
        if path.stem not in request.groups:
            path.unlink()
    levels = [level for level in ("б", "м") if any((attestation_dir / level).glob("*.xml"))]
    department = DepartmentInfo(request.department, request.faculty_title, "И. О. Фамилия", tuple(sorted(request.groups)))
    period = ritm_fetch.AttestationPeriod(2026, True, 1)
    return make_report.build_all(None, period, workdir, levels, log=log, department=department, pdf=False)


def group_job(tmp_path, groups=("ИТС(о)-11",)):
    return finished_job(tmp_path, build=offline_group_reports, codes=["ИТС(о)"], groups=list(groups))


def detailed_report_text(client, job) -> str:
    detailed = download(job, "Подробный отчёт: бакалавриат")
    document = docx.Document(io.BytesIO(client.get(detailed["href"]).content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_cannot_title_group_report_as_whole_department(tmp_path):
    _, job = group_job(tmp_path)
    assert job["markdown"].startswith("# Кафедра ИТРС, группа ИТС(о)-11:"), (
        f"Report on one group is titled as the whole department: {job['markdown'].splitlines()[0]}"
    )


def test_cannot_show_department_rank_in_group_report(tmp_path):
    _, job = group_job(tmp_path)
    assert "место кафедры" not in job["markdown"], "Report on chosen groups shows the department's rank"


def test_cannot_keep_department_section_in_group_report(tmp_path):
    client, job = group_job(tmp_path)
    assert "Место кафедры среди кафедр вуза" not in detailed_report_text(client, job), (
        "Detailed report on chosen groups describes the whole department"
    )


def test_cannot_name_level_instead_of_groups_in_short_report(tmp_path):
    client, job = group_job(tmp_path)
    assert "студентов группы ИТС(о)-11" in short_report_text(client, job), (
        "Short report on a chosen group does not name the group"
    )


def test_cannot_put_unchosen_group_in_report(tmp_path):
    _, job = group_job(tmp_path)
    assert "ИСТ(о)-21" not in job["markdown"], "A group that was not chosen got into the report"
